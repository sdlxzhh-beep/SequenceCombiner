"""Linux regression runner. Requires GCC/binutils and Python openpyxl.
Runs the actual PE with Win32 API test doubles; not a Windows/Excel UI test.
Usage: python3 src/verify.py [output-directory]
"""
from pathlib import Path
import csv
import itertools
import json
import os
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / 'verification').resolve()
OUT.mkdir(parents=True, exist_ok=True)
HARNESS = OUT / 'test_native'
subprocess.run(['gcc', '-O2', '-rdynamic', str(ROOT / 'src/test_native.c'), '-ldl', '-o', str(HARNESS)], check=True)
subprocess.run(['gcc', '-O1', '-g', '-fsanitize=address,undefined', str(ROOT / 'src/test_core.c'), '-o', str(OUT / 'test_core')], check=True)
subprocess.run([str(OUT / 'test_core'), str(OUT / 'core.xlsx')], env={**os.environ, 'ASAN_OPTIONS': 'detect_leaks=0'}, check=True)

results = []
def run(label, mode, first=None, second=None, before=None):
    folder = OUT / label
    folder.mkdir()  # Refuse stale results: each verification uses a new directory.
    if before:
        before(folder)
    args = [str(HARNESS), str(ROOT / 'SequenceCombiner.exe'), str(first or ROOT / 'examples/Group1.txt'), str(second or ROOT / 'examples/Group2.txt'), mode]
    p = subprocess.run(args, cwd=folder, capture_output=True, text=True)
    (folder / 'execution.log').write_text(p.stdout + p.stderr)
    assert p.returncode == 0, (label, p.stdout, p.stderr)
    results.append(label)
    return folder

fixtures = OUT / 'fixtures'
fixtures.mkdir(exist_ok=True)
expected = fixtures / 'expected.txt'
expected.write_text('AC\r\nTG\n', newline='')
valid = {'utf8': b'AC\r\nTG\n', 'utf8_bom': b'\xef\xbb\xbfAC\r\nTG\n',
         'utf16le': b'\xff\xfe' + 'AC\r\nTG\n'.encode('utf-16-le'),
         'utf16be': b'\xfe\xff' + 'AC\r\nTG\n'.encode('utf-16-be')}
for name, data in valid.items():
    file = fixtures / (name + '.txt'); file.write_bytes(data)
    run('import_' + name, 'import_ok', file, expected)
invalid = {'nul_utf8': b'AC\x00\nTT', 'nul_utf16': b'\xff\xfe' + 'AC\x00TT'.encode('utf-16-le'),
           'odd_le': b'\xff\xfeA\x00C', 'odd_be': b'\xfe\xff\x00AC',
           'surrogate': b'\xff\xfe\x00\xd8', 'bad_utf8': b'AC\n\xff', 'xlsx': b'PK\x03\x04\x00'}
for name, data in invalid.items():
    file = fixtures / (name + '.txt'); file.write_bytes(data)
    run('reject_' + name, 'import_bad', file, expected)

run('names', 'names')
run('validation', 'validation')
f = run('extension', 'extension'); assert not (f / 'wrong.xlsx').exists()
sentinel = b'KEEP EXISTING FILE\n'
def existing(folder):
    (folder / 'native_result.txt').write_bytes(sentinel)
for label in ('failure', 'cancel'):
    f = run(label, label, before=existing)
    assert (f / 'native_result.txt').read_bytes() == sentinel
    assert not list(f.glob('*.partial')) and not list(f.glob('*.xlsx'))
def stale(folder):
    existing(folder); (folder / 'native_result.txt.partial').write_bytes(sentinel)
f = run('stale', 'stale', before=stale)
assert (f / 'native_result.txt').read_bytes() == sentinel
assert (f / 'native_result.txt.partial').read_bytes() == sentinel

alphabet = 'ACGTRYSWKMBDHVN'
mapping = dict(zip(alphabet, 'TGCAYRSWMKVHDBN'))
def rc(seq): return ''.join(mapping[c] for c in reversed(seq))
def check_txt(path, sequences, names, reverse_flags):
    with path.open(encoding='utf-16', newline='') as f:
        reader = csv.reader(f, delimiter='\t')
        assert next(reader) == names
        count = 0
        for values in itertools.product(*sequences):
            expected_row = []
            for value, reverse in zip(values, reverse_flags):
                expected_row.append(value)
                if reverse: expected_row.append(rc(value))
            assert next(reader) == expected_row, count
            count += 1
        assert next(reader, None) is None
    assert path.read_bytes().startswith(b'\xff\xfe')
    return count
def check_xlsx(path, sequences, names, reverse_flags):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                # Streaming XML parser: do not retain hundreds of thousands of rows.
                with z.open(name) as stream:
                    for _, elem in ET.iterparse(stream): elem.clear()
    wb = load_workbook(path, read_only=True)
    rows = wb.active.iter_rows(values_only=True)
    assert next(rows) == tuple(names)
    count = 0
    for values in itertools.product(*sequences):
        expected_row = []
        for value, reverse in zip(values, reverse_flags):
            expected_row.append(value)
            if reverse: expected_row.append(rc(value))
        assert next(rows) == tuple(expected_row), count
        count += 1
    assert next(rows, None) is None
    wb.close()
    return count

sequences = [list(dict.fromkeys((ROOT / 'examples' / name).read_text(encoding='utf-8-sig').split())) for name in ('Group1.txt', 'Group2.txt')]
names = ['Group1', 'Group2', 'Group2_RC']
f = run('full', 'full')
assert check_txt(f / 'native_result.txt', sequences, names, [False, True]) == 246012
for name in ('native_result_Excel.xlsx', 'native_result.xlsx'):
    assert check_xlsx(f / name, sequences, names, [False, True]) == 246012
assert not list(f.glob('*.partial'))

f = run('companion_failure', 'companion_failure')
assert check_txt(f / 'native_result.txt', sequences, names, [False, True]) == 246012
assert not list(f.glob('*.xlsx')) and not list(f.glob('*.partial'))
def collision(folder): (folder / 'native_result_Excel.xlsx').write_bytes(sentinel)
f = run('collision', 'collision', before=collision)
assert (f / 'native_result_Excel.xlsx').read_bytes() == sentinel
assert (f / 'native_result_Excel_2.xlsx').is_file()

small = fixtures / 'small.txt'; small.write_text('AC\nTG\n')
f = run('unicode', 'unicode', small, small)
unicode_names = ['A&_x0041_<"', 'A&_x0041_<"_RC', '组二😀', '组二😀_RC']
assert check_txt(f / 'native_result.txt', [['AC', 'TG']] * 2, unicode_names, [True, True]) == 4
# Inspect encoded SpreadsheetML text; openpyxl does not unescape inline strings.
with zipfile.ZipFile(f / 'native_result_Excel.xlsx') as z:
    xml = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    heads = xml.find('s:sheetData/s:row', ns)
    decoded = [c.find('s:is/s:t', ns).text.replace('_x005F_', '_') for c in heads]
    assert decoded == unicode_names
    assert z.testzip() is None

f = run('columns64', 'columns64')
names64 = [name for i in range(1, 33) for name in (f'Group{i}', f'Group{i}_RC')]
assert check_xlsx(f / 'columns64.xlsx', [['AC']] * 32, names64, [True] * 32) == 1

long_seq = fixtures / 'long.txt'; long_seq.write_text('A' * 32768 + '\n')
one = fixtures / 'one.txt'; one.write_text('AC\n')
f = run('longcell', 'limit', long_seq, one)
assert not list(f.glob('*.xlsx'))
csv.field_size_limit(1000000)
assert check_txt(f / 'native_result.txt', [['A' * 32768], ['AC']], names, [False, True]) == 1

report = {'version': '1.2.1', 'scenarios': results, 'scenario_count': len(results), 'full_combinations': 246012,
          'checks': ['Core AddressSanitizer/UndefinedBehaviorSanitizer', 'PE machine-code execution with Win32 doubles',
                     'UTF-16 TXT full row comparison', 'Two XLSX exports: ZIP CRC/XML parsing/openpyxl full row comparison',
                     'Import encodings and corruption rejection', 'No replacement on cancellation or write failure',
                     'Existing companion and stale temporary file preservation', '64 columns', 'Excel cell limit'],
          'limitation': 'Not tested on a real Windows desktop or Microsoft Excel.'}
(OUT / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False, indent=2))
